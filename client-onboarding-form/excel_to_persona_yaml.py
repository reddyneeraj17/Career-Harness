#!/usr/bin/env python3
"""Convert a filled Client Onboarding Form (Excel v2) into a client persona YAML.

Usage:
    python3 excel_to_persona_yaml.py <form.xlsx> [--out profiles/<client_id>.yaml]

Reads the v2 sheets (Start Here / 1 About You / ... / 10 Sign-off), validates
required fields, and writes profiles/<client_id>.yaml conforming to
harness-kit/templates/client-persona.schema.yaml. Never writes passwords —
the form must not contain any; the converter refuses values under
password-like labels.
"""
import re
import sys
import os
from datetime import date, datetime

import yaml
from openpyxl import load_workbook

PROFILES_DIR = os.path.expanduser("~/workspace/profiles")

FORBIDDEN_LABEL_HINTS = ("password", "passwd", "pwd", "one-time", "otp", "totp",
                          "card number", "cvv", "bank account", "ssn",
                          "social security")


def s(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    t = str(v).strip()
    return t


def to_int(v, default=0):
    try:
        return int(float(str(v).strip()))
    except (ValueError, TypeError):
        return default


def to_bool(v):
    t = str(v).strip().lower()
    if t in ("yes", "true", "y", "1"):
        return True
    if t in ("no", "false", "n", "0"):
        return False
    return None


def split_list(v):
    t = s(v)
    if not t:
        return []
    parts = re.split(r"[,;\n]+", t)
    return [p.strip() for p in parts if p.strip()]


def slug(name):
    t = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return t or "client-unknown"


# Canonical employment-lane slugs used by eligibility_judge (`profile.role_types`).
LANE_SLUGS = ("full_time", "part_time", "w2_contract", "c2c_contract", "internship")

def normalize_employment_types(v):
    """Map the form's display labels ('Full-time + W2 + C2C') to canonical
    lane slugs. Keyword-based so both dropdown values and free-typed answers
    work. Returns [] when nothing recognizable is present."""
    t = s(v).lower()
    if not t:
        return []
    if "all" in t and "type" in t:
        return list(LANE_SLUGS)
    lanes = []
    if "full-time" in t or "full time" in t:
        lanes.append("full_time")
    if "part-time" in t or "part time" in t:
        lanes.append("part_time")
    if re.search(r"\bw2\b", t):
        lanes.append("w2_contract")
    if "c2c" in t:
        lanes.append("c2c_contract")
    if "intern" in t:
        lanes.append("internship")
    if not lanes and "contract" in t:
        # "Contract only" with no finer detail -> both contract lanes.
        lanes = ["w2_contract", "c2c_contract"]
    # Preserve canonical order, drop dupes.
    return [l for l in LANE_SLUGS if l in lanes]


def normalize_lane(v):
    """Single lane slug for `preferred_lane` ('W2 contract' -> 'w2_contract')."""
    lanes = normalize_employment_types(v)
    return lanes[0] if lanes else ""


def rows_of(wb, name):
    ws = wb[name]
    return list(ws.iter_rows(values_only=True))


def field_map(rows):
    """label -> value for label/value sheets (col A = label, col B = value)."""
    d = {}
    for r in rows:
        if not r or r[0] is None:
            continue
        label = str(r[0]).strip()
        if not label:
            continue
        low = label.lower()
        # Only short, prompt-style labels can be secret requests; long
        # attestation sentences (e.g. the sign-off "I have not written any
        # password...") must not trip this guard.
        if len(label) < 60 and any(h in low for h in FORBIDDEN_LABEL_HINTS):
            raise ValueError(f"Refusing to convert: form asks for a secret in field '{label}'. "
                             "Remove secrets from the form.")
        d[label] = s(r[1]) if len(r) > 1 else ""
    return d


def pick(d, *needles):
    best, best_len = "", None
    for label, val in d.items():
        if all(n.lower() in label.lower() for n in needles):
            # Prefer the shortest matching label: exact field labels
            # ("Your name") beat instruction sentences that mention them.
            if best_len is None or len(label) < best_len:
                best, best_len = val, len(label)
    return best


def header_index(rows, *needles):
    # Several candidate rows can match (e.g. a rules sentence mentioning
    # "skills" and "years" above the real header). Prefer the match whose
    # cells are shortest overall — real headers are terse, sentences are not.
    best, best_len = None, None
    for i, r in enumerate(rows):
        vals = [str(v).strip() if v is not None else "" for v in r]
        low = [v.lower() for v in vals]
        if all(any(n in v for v in low) for n in needles):
            total = sum(len(v) for v in vals)
            if best_len is None or total < best_len:
                best, best_len = i, total
    return best


def _is_int_text(t):
    try:
        int(str(t).strip())
        return True
    except (ValueError, TypeError):
        return False


def table_after(rows, hidx, ncols, stop_words=(), require_int_first=False):
    out = []
    for r in rows[hidx + 1:]:
        vals = [s(v) for v in (list(r) + [""] * ncols)[:ncols]]
        a0 = vals[0].lower()
        if stop_words and any(w in a0 for w in stop_words):
            break
        if not any(vals):
            continue
        # Note/instruction rows below a table have no numeric index —
        # real data rows always do (# / Priority column).
        if require_int_first and not _is_int_text(vals[0]):
            # ...except the essay-policy row in Screening Answers.
            if not ("essay" in vals[1].lower() and not vals[0]):
                continue
        out.append(vals)
    return out


def convert(form_path):
    wb = load_workbook(form_path, data_only=True)

    about = field_map(rows_of(wb, "1 About You"))
    auth = field_map(rows_of(wb, "2 Work Authorization"))
    prefs = field_map(rows_of(wb, "6 Preferences"))
    auto = field_map(rows_of(wb, "9 Automation Settings"))
    sign = field_map(rows_of(wb, "10 Sign-off"))
    acct = field_map(rows_of(wb, "8 Accounts & Connections"))

    # ---- 3 Resumes ----
    r_rows = rows_of(wb, "3 Resumes")
    rh = header_index(r_rows, "#", "resume label", "file name")
    resumes = []
    for vals in table_after(r_rows, rh, 8, require_int_first=True):
        # #, label, file name, role family, industry tags, last updated, primary, notes
        if not vals[1] and not vals[2]:
            continue
        resumes.append({
            "label": vals[1], "file_name": vals[2], "role_family": vals[3],
            "industry_tags": split_list(vals[4]), "last_updated": vals[5],
            "primary": bool(to_bool(vals[6])), "notes": vals[7],
        })

    # ---- 4 Experience Matrix ----
    e_rows = rows_of(wb, "4 Experience Matrix")
    eh = header_index(e_rows, "skill", "years")
    years_matrix = []
    # Stop at the "Employment history" section below the skills table — its
    # numbered rows would otherwise pollute years_matrix with 0-year entries.
    for vals in table_after(e_rows, eh, 5, require_int_first=True, stop_words=("employment history",)):
        if not vals[1]:
            continue
        years_matrix.append({
            "skill": vals[1], "category": vals[2],
            "years": to_int(vals[3]), "where_used": vals[4],
        })

    # ---- 5 Target Roles ----
    t_rows = rows_of(wb, "5 Target Roles")
    th = header_index(t_rows, "priority", "job title")
    target_roles = []
    for vals in table_after(t_rows, th, 7, require_int_first=True):
        if not vals[1]:
            continue
        target_roles.append({
            "priority": to_int(vals[0], 99), "title": vals[1],
            "role_family": vals[2], "seniority": vals[3],
            "industries": split_list(vals[4]), "keywords": vals[5],
            "avoid_titles": vals[6],
        })
    target_roles.sort(key=lambda r: r["priority"])

    # ---- 5b Company targeting (label/value rows below the roles table) ----
    corp = field_map(t_rows)
    company_targeting = {
        "tier_preference": pick(corp, "company size", "tier preference"),
        "industries_to_avoid": split_list(pick(corp, "industries to avoid")),
        "never_apply_companies": split_list(pick(corp, "never apply")),
        "dream_companies": split_list(pick(corp, "dream companies")),
    }

    # ---- 7 Screening Answers ----
    sc_rows = rows_of(wb, "7 Screening Answers")
    sch = header_index(sc_rows, "#", "question", "your standard answer")
    answers, essay_policy = [], ""
    for vals in table_after(sc_rows, sch, 4, require_int_first=True):
        if not vals[1] and not vals[2]:
            continue
        if not vals[0] and "essay" in vals[1].lower():
            essay_policy = vals[2]
            continue
        if not vals[1]:
            continue
        answers.append({"question": vals[1], "answer": vals[2], "notes": vals[3]})

    # ---- 8 Portals table ----
    a_rows = rows_of(wb, "8 Accounts & Connections")
    ph = header_index(a_rows, "platform", "status", "email")
    portals = []
    for vals in table_after(a_rows, ph, 6,
                            stop_words=("mailbox", "muse account", "portal accounts we may")):
        if not vals[0]:
            continue
        portals.append({
            "platform": vals[0], "status": vals[1],
            "email_or_username": vals[2], "profile_url": vals[3],
            "two_fa": vals[4], "notes": vals[5],
        })

    # ---- 10 Sign-off confirmations ----
    confirmations = []
    for label, val in sign.items():
        ll = label.lower()
        if ll in ("your name", "date (yyyy-mm-dd)", "date"):
            continue
        if val.lower() in ("yes", "no") and len(label) > 12:
            confirmations.append({"item": label, "answer": val})

    full_name = pick(about, "full name")
    persona = {
        "client_id": slug(full_name),
        "version": 1,
        "source_form": os.path.basename(form_path),
        "identity": {
            "full_name": full_name,
            "preferred_name": pick(about, "preferred first name"),
            "phone": pick(about, "phone"),
            "email": pick(about, "personal email"),
            "city": pick(about, "city") if pick(about, "city") else "",
            "state": pick(about, "state"),
            "timezone": pick(about, "timezone"),
            "street_address": pick(about, "street address"),
            "linkedin_url": pick(about, "linkedin profile url"),
            "github_url": pick(about, "github / portfolio url"),
            "website": pick(about, "personal website"),
            "current_title": pick(about, "most recent job title"),
            "current_employer": pick(about, "most recent employer"),
            "employment_status": pick(about, "employment status"),
            "total_years_experience": to_int(pick(about, "total years")),
            "earliest_start_date": pick(about, "earliest start date"),
            "notice_period": pick(about, "notice period"),
            "headline": pick(about, "headline we may use"),
        },
        "work_auth": {
            "status": pick(auth, "current work authorization"),
            "visa_expiry": pick(auth, "visa expiry"),
            "sponsorship_required": pick(auth, "need sponsorship"),
            "sponsorship_sentence": pick(auth, "standard sentence"),
            "h1b_gate": {"strict": "strict", "prefer": "prefer", "ignore": "ignore"}.get(
                next((k for k in ("strict", "prefer", "ignore")
                      if k in pick(auth, "sponsorship-history filter").lower()), ""), ""),
            "apply_without_sponsorship_mention": pick(auth, "do not mention sponsorship"),
            "skip_no_sponsorship_postings": pick(auth, "skip postings that say"),
            "notes": pick(auth, "immigration situation"),
        },
        "resumes": resumes,
        "years_matrix": years_matrix,
        "target_roles": target_roles,
        "company_targeting": company_targeting,
        "preferences": {
            "work_mode": pick(prefs, "work mode"),
            "open_to_relocation": pick(prefs, "open to relocation"),
            "preferred_metros": split_list(pick(prefs, "preferred cities")),
            "us_only": to_bool(pick(prefs, "us only?")),
            "employment_types": normalize_employment_types(
                pick(prefs, "employment types you will accept")),
            "preferred_lane": normalize_lane(
                pick(prefs, "preferred employment lane")),
            "agencies_ok": pick(prefs, "staffing agencies"),
            "max_travel": pick(prefs, "maximum travel"),
            "min_base_salary_usd": to_int(pick(prefs, "minimum base salary"), None)
            if pick(prefs, "minimum base salary") else None,
            "desired_salary_answer": pick(prefs, "desired salary"),
            "salary_number_fallback": pick(prefs, "requires a number"),
            "interview_windows": pick(prefs, "interview availability"),
        },
        "screening": {"answers": answers, "essay_policy": essay_policy},
        "accounts": {
            "portals": portals,
            "mailbox_scan": pick(acct, "mailbox we should scan"),
            "mailbox_address": pick(acct, "address of that mailbox"),
            "mailbox_created": pick(acct, "already created"),
            "mailbox_never_scan": pick(acct, "must never scan"),
            "muse": {
                "has_account": pick(acct, "have a muse account"),
                "email": pick(acct, "existing muse email"),
                "phone": pick(acct, "muse verification"),
            },
            "portal_creation": {
                "may_create": to_bool(pick(acct, "may we create new accounts")),
                "standard_password": "",
            },
        },
        "automation": {
            "max_per_day": to_int(pick(auto, "maximum applications per day"), 20),
            "max_per_run": to_int(pick(auto, "maximum applications per run"), 10),
            "hours": pick(auto, "hours we may operate"),
            "linkedin_actions_per_hour": to_int(pick(auto, "linkedin actions per hour"), 20),
            "auto_send": {
                "routine_email_replies": pick(auto, "routine email replies"),
                "interview_scheduling": pick(auto, "interview scheduling"),
                "resume_requests": pick(auto, "resume requests by email"),
                "followup_nudges": pick(auto, "follow-up nudges"),
                "linkedin_replies": pick(auto, "linkedin replies to recruiters"),
                "linkedin_outreach": pick(auto, "linkedin outreach to hiring managers"),
                "decline_contract_pitches": pick(auto, "declining contract"),
            },
            "notifications": {
                "frequency": pick(auto, "how often do you want updates"),
                "channel": pick(auto, "best channel"),
                "best_times": pick(auto, "best times to reach you"),
                "immediate_on_interview": pick(auto, "tell me immediately"),
            },
            "do_not_contact": split_list(pick(auto, "never reply to")),
            "tone_notes": pick(auto, "tone or style"),
            "hold_policy": pick(auto, "when the system is unsure", "hold policy"),
        },
        "signoff": {
            "confirmations": confirmations,
            "name": pick(sign, "your name"),
            "date": pick(sign, "date"),
        },
    }

    # City is its own label; the generic pick above may catch "Preferred cities" — fix:
    for label, val in about.items():
        if label.strip().lower() == "city":
            persona["identity"]["city"] = val
            break

    warnings = []
    for key, path in (("identity.full_name", ["identity", "full_name"]),
                      ("identity.email", ["identity", "email"]),
                      ("work_auth.status", ["work_auth", "status"]),
                      ("signoff.name", ["signoff", "name"]),
                      ("signoff.date", ["signoff", "date"])):
        node = persona
        for p in path:
            node = node.get(p, {})
        if not node:
            warnings.append(f"missing required: {key}")
    if not resumes:
        warnings.append("missing required: at least one resume row")
    if not target_roles:
        warnings.append("missing required: at least one target role row")
    return persona, warnings


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    form_path = sys.argv[1]
    persona, warnings = convert(form_path)
    out = None
    for a in sys.argv[2:]:
        if a == "--out" :
            idx = sys.argv.index(a)
            out = sys.argv[idx + 1]
    if out is None:
        os.makedirs(PROFILES_DIR, exist_ok=True)
        out = os.path.join(PROFILES_DIR, persona["client_id"] + ".yaml")

    header = (
        "# Client persona — generated from the onboarding form.\n"
        "# Source of truth for all agents working this client's job search.\n"
        "# NEVER add passwords, one-time codes, or payment details here.\n"
    )
    with open(out, "w") as f:
        f.write(header)
        yaml.safe_dump(persona, f, sort_keys=False, allow_unicode=True, width=100)
    print("wrote:", out)
    if warnings:
        print("WARNINGS (operator should resolve with the client):")
        for w in warnings:
            print(" -", w)
    else:
        print("validation: required fields present")


if __name__ == "__main__":
    main()
